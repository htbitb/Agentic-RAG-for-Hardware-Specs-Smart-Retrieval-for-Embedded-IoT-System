#!/usr/bin/env python3
"""Dựng PPTX tiếng Việt cho Báo cáo tiến độ giữa kỳ D2 (Hardware Co-Design Agent v2).
Chạy: /tmp/interim-publish-venv/bin/python build_slides.py
Xuất PDF: soffice --headless --convert-to pdf --outdir . interim-slides.pptx
"""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'interim-slides.pptx'
prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
P = dict(white='FFFFFF', ink='18344B', blue='164B78', teal='16756A', amber='925D16', muted='516273', pale='EDF3F8', line='CBD7E1', light='F6F8FA')
FONT = 'Arial'

def rgb(c): return RGBColor.from_string(P.get(c, c))
def rect(s,x,y,w,h,fill='pale',line=None,radius=False):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
    if line: sh.line.color.rgb=rgb(line); sh.line.width=Pt(1)
    else: sh.line.fill.background()
    if radius: sh.adjustments[0]=0.08
    return sh

def txt(s,x,y,w,h,text,size=21,bold=False,color='ink',align=None,font=FONT):
    sh=s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf=sh.text_frame; tf.clear(); tf.word_wrap=True
    tf.margin_left=0; tf.margin_right=0; tf.margin_top=0; tf.margin_bottom=0
    for i,line in enumerate(text.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.font.name=font; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=rgb(color)
        p.space_after=Pt(4); p.space_before=Pt(0); p.line_spacing=1.08
        if align is not None: p.alignment=align
    return sh

def line(s,x1,y1,x2,y2,color='blue',width=1.5):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    sh.line.color.rgb=rgb(color); sh.line.width=Pt(width)
    return sh

def note(s,t): s.notes_slide.notes_text_frame.text=t

def base(title,sub='',source='',time='',backup=False):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb('white')
    rect(s,0,0,.16,7.5,'blue')
    txt(s,.6,.32,12.1,.56,title,28,True,'blue')
    if sub: txt(s,.62,.92,12.0,.42,sub,15.5,False,'muted')
    line(s,.62,6.96,12.72,6.96,'line',.8)
    txt(s,.62,7.08,10.8,.23,source or 'CO5151 • Báo cáo D2 • Agent v2 (S3@P5) • 24/10/2026',9.5,False,'muted')
    idx=len(prs.slides)
    txt(s,11.65,7.06,1.05,.25,('Phụ lục ' if backup else '')+str(idx),10,False,'muted',PP_ALIGN.RIGHT)
    return s

def table(s,x,y,width,headers,rows,widths=None,size=17,row_h=.48,highlight=None,hdr_h=None):
    t=s.shapes.add_table(len(rows)+1,len(headers), Inches(x), Inches(y), Inches(width), Inches((hdr_h or row_h)+row_h*len(rows))).table
    if widths:
        for col,w in zip(t.columns,widths): col.width=Inches(w)
    for r,data in enumerate([headers]+rows):
        t.rows[r].height=Inches((hdr_h or row_h) if r==0 else row_h)
        for c,val in enumerate(data):
            cell=t.cell(r,c); cell.text=str(val); cell.margin_left=Inches(.10); cell.margin_right=Inches(.08); cell.margin_top=Inches(.05); cell.margin_bottom=Inches(.04)
            cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid(); cell.fill.fore_color.rgb=rgb('blue' if r==0 else ('pale' if r==highlight else ('light' if r%2 else 'white')))
            for p in cell.text_frame.paragraphs:
                p.font.name=FONT; p.font.size=Pt(size); p.font.bold=(r==0 or r==highlight)
                p.font.color.rgb=rgb('white' if r==0 else 'ink'); p.space_after=Pt(0); p.line_spacing=1.04
    return t

def band(s,text,y=6.15,color='teal',size=20):
    rect(s,.62,y,12.05,.55,'pale')
    txt(s,.80,y+.10,11.7,.39,text,size,True,color)

# 1 — Trang bìa (15 s)
s=base('Báo cáo tiến độ giữa kỳ (D2)',source='CO5151 • HK261 • GVHD: TS. Lê Xuân Bách • Application Track')
s.shapes.add_picture(str(ROOT/'assets/logo.png'), Inches(11.3), Inches(1.65),width=Inches(1.12))
txt(s,.72,1.62,10.1,1.30,'Agentic RAG\nfor Hardware Specs',39,True,'blue')
txt(s,.75,3.18,11.3,.85,'Từ tra cứu thông số đến Hardware Co-Design Agent v2\ncho hệ thống nhúng IoT',25,False)
line(s,.75,4.36,10.8,4.36,'line')
txt(s,.75,4.62,11.6,1.08,'Hà Thanh Bình — 2470732 (33,33%)\nHuỳnh Đức Nhâm — 2570276 (33,33%)\nLa Quang Huy — 2470736 (33,33%)',18)
txt(s,.75,6.12,11.5,.40,'Phạm vi đã cài đặt: Agent v2 — P0–P5 (S3@P5) • Tiến độ: đúng tiến độ',18,True,'teal')
note(s,'15 giây. Giới thiệu đề tài và 3 thành viên. Phiên bản báo cáo là Agent v2 (cấu hình đánh giá S3@P5), đã hoàn thành các giai đoạn P0–P5. Bài trình bày gồm 10 phút thuyết trình kèm live demo và 5 phút hỏi đáp.')

# 2 — Thay đổi so với proposal (35 s)
s=base('1. Thay đổi so với proposal (Nâng cấp v1 → v2)','Nguồn kế hoạch nâng cấp: plan/AgenticUpgradePlan.md (§0, §1, §8)',source='D1: plan/proposal.pdf • Kế hoạch nâng cấp v2: plan/AgenticUpgradePlan.md')
table(s,.65,1.55,12.0,['Cam kết D1 / Bản v1','Trạng thái v2 đã chạy (P0–P5)','Kế hoạch hoàn thiện (P6–P9)'],[
['Câu hỏi đơn → tra cứu chỉ đọc','Mục tiêu thiết kế → BoardState + Sandbox','Policy Engine cấp hành động (P6)'],
['Verifier đếm số khớp văn bản','Gate G1–G5 + vòng tự sửa có kiểu lỗi','Project & Lesson Memory (P7)'],
['Duyệt file ở cuối, không chạy tiếp','Interrupt / Resume qua SQLite checkpoint','Giao diện Web UI v2 Panels (P8)'],
['40+ task; baseline & ablation','132 task (C1–C8); đo 42 task dev × 3 seed','Chạy tập test niêm phong, C9, A1–A7 (P9)'],
],widths=[3.55,4.35,4.10],size=17.5,row_h=.82)
band(s,'Thay đổi cốt lõi: cho agent môi trường BoardState có trạng thái và phản hồi từ Simulator / DRC / Compiler.',6.05,size=18.5)
note(s,'35 giây. Trình bày sự nâng cấp từ Agentic RAG chỉ đọc ở v1 lên Hardware Co-Design Agent v2. Mọi cam kết giữa kỳ về dữ liệu PDF/SVD, MCP servers, vòng điều khiển R1, bộ nhớ R3 và khung đánh giá R5 đều đã vận hành trên tập dev.')

# 3 — Kiến trúc đã cài đặt (45 s) — nhúng sơ đồ draw.io
s=base('2. Kiến trúc hệ thống Agent v2 đã cài đặt (R1)','Một LLM Agent điều phối bằng State Machine có Checkpoint • 4 MCP Servers kết nối inproc',source='docs/architecture.md • agent/agent/runtime/ • agent/agent/gate/')
s.shapes.add_picture(str(ROOT/'assets/architecture_drawio.png'), Inches(0.65), Inches(1.42), width=Inches(12.0))
note(s,'45 giây. Sơ đồ kiến trúc v2 vẽ theo chuẩn draw.io gồm 4 khối: (1) Giao tiếp CLI & Run API; (2) Agent Runtime v2 với INTAKE, PLAN, ACT<->OBSERVE, CHECK & FINISH_GATE G1–G5, và INTERRUPTED<->RESUME; (3) ToolHub quản lý 32 công cụ theo 5 mức quyền; (4) Tầng dữ liệu thật và SQLite. Các hộp viền nét đứt màu cam là phần kế hoạch P6–P8.')

# 4 — Công cụ & Bộ nhớ (35 s)
s=base('3–4. Phân quyền 32 công cụ (R2) và Bộ nhớ & Tri thức (R3)','Cưỡng chế quyền ở ToolHub • Không phơi bày shell hay thực thi mã nhị phân do LLM sinh ra',source='agent/agent/tools.py • runtime/hitl.py • ingestion/ • docs/architecture.md')
table(s,.65,1.58,7.35,['Server / Nhóm (32 tools)','Mức quyền & Cơ chế chặn ở mã nguồn'],[
['text-manual (3)\nhw-knowledge (8)','Read-only: chỉ đọc FTS5/vector và SQLite tham số hoá; không nhận SQL hay đường dẫn tuỳ ý'],
['board-state (13:\n4 đọc, 9 ghi)','State-mutating: event-sourced theo phiên bản, kiểm tra xung đột chân/áp/clock, hỗ trợ rollback'],
['hw-sandbox (4)','Sandboxed-exec: solver, parser C, simulator thanh ghi, clang -fsyntax-only (không chạy binary)'],
['local-control (3)\nlocal-guarded (1)','Control & Irreversible-write: giới hạn ngân sách; write_config_file bắt buộc qua cổng Approval'],
],widths=[2.25,5.10],row_h=.86,size=16)
txt(s,8.25,1.62,4.4,.45,'Bộ nhớ & Nguồn tri thức',22,True,'blue')
txt(s,8.25,2.10,4.4,2.10,'• Ngắn hạn: giữ 12 tool messages mới nhất, rút gọn message cũ còn 1.500 ký tự.\n• Bền vững (SQLite): Checkpoint từng bước, Session History và BoardState.\n• Kế hoạch P7: Lesson Memory lưu bài học sau khi simulator xác nhận.',16.5)
rect(s,8.25,4.35,4.4,1.52,'pale',line='teal',radius=True)
txt(s,8.42,4.46,4.1,1.38,'Quy mô dữ liệu thật đã nạp:\n• 6 PDF (4.520 trang) → 9.583 chunk\n• 3 SVD → 210 periph, 6.338 thanh ghi\n• 27.186 bitfield, 189 chân / 1.045 AF',15.5,True,'teal')
band(s,'100% giới hạn điện (65) và xung nhịp (39) đều gắn số trang và trích dẫn khớp PDF gốc.',6.08,size=18.5)
note(s,'35 giây. Tóm tắt 32 công cụ chia theo 5 mức quyền và quy mô kho tri thức thật từ 6 PDF hãng và 3 file CMSIS-SVD.')

# 5 — Thiết lập đánh giá (35 s)
s=base('5. Bộ đánh giá xác định trên 132 task (R5)','Model: ag/gemini-3.8-flash-medium (temperature = 0,0) • Chấm điểm tự động 100% xác định',source='eval/README.md • eval/results/dev_baseline_v2/manifest.json')
table(s,.65,1.58,12.0,['Quy mô & Phân bổ tập Task','Bộ chấm điểm xác định (Deterministic Graders)','Thiết lập đối sánh công bằng'],[
['132 task gốc (C1–C8):\n• 42 task dev (đã đo)\n• 90 task test (niêm phong)\n• Mục tiêu D3: ~150 task (+C9)','• C1, C2: exact / keyfacts + citation\n• C3: functional_sim (simulator đóng băng)\n• C4: final_state + conflict_detected\n• C5–C8: checklist, safety, attack, abstain','• 5 hệ: S0, S1, S2, S2D, S3@P5\n• Cùng LLM, cùng corpus gốc\n• 3 seed (0, 1, 2) / task\n• CI 95% bootstrap (10.000 mẫu)'],
],widths=[3.65,4.45,3.90],size=17,row_h=3.15,hdr_h=.52)
band(s,'Chống tự chấm: Grader dùng bản simulator đóng băng độc lập trong eval/vendor/hwmodel/.',6.05,size=19)
note(s,'35 giây. Trình bày bộ 132 task thuộc 8 nhóm C1–C8 và cơ chế chấm điểm xác định bằng simulator độc lập.')

# 6 — Kết quả chính (50 s)
s=base('5. Kết quả thực nghiệm trên tập dev (42 task × 3 seed)','TSR (%) theo từng nhóm task và TSR tổng kèm khoảng tin cậy 95% bootstrap',source='eval/results/dev_baseline_v2/report.md • 774 run đã chấm')
table(s,.65,1.58,12.0,['Hệ thống','C1','C2','C3','C4','C5','C6','C7','C8','TSR tổng [CI 95%]'],[
['S0 Direct LLM','79,2','77,8','77,8','36,7','41,7','66,7','93,3','100','71,8% [60,3–82,9]'],
['S1 Native RAG','25,0','50,0','0,0','0,0','0,0','0,0','100','100','31,0% [16,7–45,2]'],
['S2 Agent v1','0,0','5,6','5,6','0,0','0,0','20,0','100','100','23,0% [11,1–35,7]'],
['S2D (v1 + dữ liệu)','100','83,3','0,0','0,0','0,0','53,3','100','100','56,3% [41,3–70,7]'],
['S3@P5 (Agent v2)','100','100','100','60,0','58,3','80,0','100','77,8','87,3% [77,8–95,2]'],
],widths=[2.65,0.85,0.85,0.85,0.85,0.85,0.85,0.85,0.85,2.60],size=16.5,row_h=.58,highlight=5)
txt(s,.80,5.28,11.8,.45,'• So với S2D: +31,0 điểm phần trăm (CI [+17,5; +45,2], McNemar p = 0,001) → đóng góp rõ rệt của kiến trúc v2.',18,True,'teal')
txt(s,.80,5.96,11.8,.45,'• So với S0: +15,5 điểm phần trăm (CI [+2,0; +28,6], p = 0,146); Grounded TSR C1+C2 đạt 85,7% so với 0,0% của S0.',18,True,'blue')
note(s,'50 giây. Nhấn mạnh bước nhảy từ S2 lên S2D (+33,3 điểm nhờ dữ liệu thật P1) và từ S2D lên S3@P5 (+31,0 điểm nhờ kiến trúc đồng thiết kế v2, đặc biệt trên C3, C4, C5).')

# 7 — Chi phí & Độ trễ (30 s)
s=base('5. Chi phí, Độ trễ và Chỉ số quá trình Agentic','Đơn giá cấu hình: 0,075 / 0,30 USD mỗi triệu token vào/ra • Token cộng dồn toàn bộ các bước trong run',source='eval/results/dev_baseline_v2/report.md • Chỉ số quá trình S3@P5')
table(s,.65,1.62,12.0,['Hệ thống','Grounded TSR','Citation Prec.','Token vào / ra','USD / run','p50 / p95 (s)'],[
['S0 Direct LLM','0,0%','0,0%','5.175 / 996','0,00069','16,7 / 90,0'],
['S1 Native RAG','35,7%','23,2%','4.723 / 20','0,00036','7,3 / 22,1'],
['S2 Agent v1','0,0%','0,0%','37.946 / 330','0,00294','30,4 / 92,9'],
['S2D (v1 + dữ liệu)','85,7%','89,9%','61.634 / 446','0,00476','30,2 / 121,5'],
['S3@P5 (Agent v2)','85,7%','86,6%','406.677 / 5.234','0,03207','35,6 / 518,0'],
],widths=[2.55,1.75,1.75,2.45,1.60,1.90],size=17.5,row_h=.58,highlight=5)
band(s,'Tự sửa lỗi vật lý: 8/9 run • Tự sửa nguồn G1: 52/53 run • Clarification P/R: 4/4 (100%) và 6/9 (66,7%).',5.95,color='teal',size=18.5)
note(s,'30 giây. Chi phí trung bình mỗi run của S3@P5 là 0,032 USD (~6,7 lần S2D) do vòng lặp nhiều bước gọi công cụ và kiểm chứng Gate.')

# 8 — Trace thực tế (35 s)
s=base('2. Trace thực tế: Phê duyệt ghi file & Chạy tiếp từ Checkpoint','Run 6d4db321f405 • Thiết kế PWM PA8/TIM1 20 kHz, duty 25% ở SYSCLK 170 MHz',source='plan/demos/D_approval/run1/summary.md • trace/6d4db321f405.jsonl')
table(s,.65,1.65,12.0,['Seq / Step','Sự kiện trong Trace (LLM thật)'],[
['seq 1–29 / step 1–10','Đặt target STM32G474, tra PA8 -> AF6 (TIM1_CH1), giải PSC=0, ARR=8499, CCR1=2125, bật BDTR.MOE'],
['seq 30 / step 11','interrupt_raised (approval): xin ghi pwm_pa8.c (1.579 byte); lưu checkpoint INTERRUPTED'],
['seq 31–33 / step 11','Tiến trình mới gọi resume(approve) -> file_written vào workspace_out/pwm_pa8.c'],
['step 13–15','Agent chạy tiếp: cập nhật requirement & plan, chạy lại run_drc đạt yêu cầu'],
['step 16','Mô phỏng đạt 20.000 Hz, duty 25%; Gate G1–G5 đều pass -> finish ACCEPTED'],
],widths=[2.95,9.05],size=18,row_h=.68)
band(s,'Chứng minh cơ chế Interrupt/Resume bền vững qua khởi động lại tiến trình và cổng kết thúc G1–G5.',6.02,color='teal',size=19)
note(s,'35 giây. Trình bày trace thực tế của phiên chạy 6d4db321f405 từ lúc lập kế hoạch, giải tham số timer, tạm dừng xin phép ghi file, khôi phục ở tiến trình mới cho tới khi vượt qua Gate G1–G5.')

# 9 — Kịch bản Live Demo (120 s)
s=base('Kịch bản Live Demo (2 phút): Thiết kế PWM & Cổng phê duyệt','Chạy trực tiếp qua CLI / Run API v2 với checkpoint SQLite',source='plan/demos/D_approval/run1/ • agent/agent/runtime/cli.py')
txt(s,.8,1.58,11.8,.82,'“Trên STM32G474 chạy 170 MHz, tạo PWM 20 kHz, duty 25% ở PA8 (TIM1_CH1)\nvà lưu mã cấu hình CMSIS vào tệp pwm_pa8.c.”',25,True,'blue')
table(s,.75,2.62,11.8,['Mốc thời gian','Nội dung quan sát trực tiếp trong buổi trình bày'],[
['0:00 – 0:25','Khởi tạo mục tiêu, lập kế hoạch (update_plan), tra cứu PA8 AF6 và tính PSC=0, ARR=8499, CCR1=2125'],
['0:25 – 1:00','Agent gọi write_config_file -> Runtime dừng ở trạng thái INTERRUPTED; kiểm tra tệp chưa được ghi'],
['1:00 – 1:35','Người dùng gửi lệnh resume(approve) -> Tệp pwm_pa8.c xuất hiện trong workspace_out/, agent chạy tiếp'],
['1:35 – 2:00','Simulator xác nhận 20 kHz / 25% duty; Gate G1–G5 pass; đọc mã C sinh ra và nhật ký kiểm toán'],
],widths=[2.45,9.35],row_h=.60,hdr_h=.46,size=16.5)
band(s,'Phương án dự phòng: phát lại từ cache (replay) hoặc video demo đã chuẩn bị sẵn nếu mạng gián đoạn.',5.98,color='teal',size=18.5)
note(s,'120 giây. Thực hiện live demo theo 4 bước rõ ràng, có sẵn checkpoint và video dự phòng để đảm bảo thời lượng 10 phút.')

# 10 — Threat Model (35 s)
s=base('6. Threat Model & Cơ chế bảo vệ đa tầng (R6)','Bảo vệ khoá API, kho dữ liệu, BoardState, thư mục workspace_out/ và an toàn phần cứng',source='agent/agent/guardrail.py • runtime/hitl.py • plan/HANDOFF_P6_P9.md')
table(s,.65,1.56,12.0,['Điểm vào & Nguy cơ','Biện pháp giảm thiểu đã cài đặt (S3@P5)','Kế hoạch nâng cấp ở P6'],[
['Câu hỏi người dùng (Direct / quá áp)','Regex INJECTION_PATTERNS & hardware_check tại INTAKE; Gate G4 ở đầu ra','Kiểm tra tham số trên từng hành động để loại bỏ over-refusal (C6-015)'],
['Tài liệu PDF / Observation (Indirect)','Khử độc đệ quy (sanitize) kết quả tool; kiểm nguồn G1; ghim SHA-256 corpus','Bổ sung Taint Tracking chặn ghi tệp khi dữ liệu đầu vào dính cờ cảnh báo'],
['Công cụ ghi tệp (write_config_file)','Allowlist tên/đuôi tệp, ≤ 20 KB, chống path traversal; bắt buộc người dùng Approve','Thêm hạn ngạch số tệp ghi mỗi phiên và khoá cấu hình khi chạy đa người dùng'],
['BoardState & Mã C do LLM sinh ra','Event log chỉ nối thêm (append-only), rollback; mã C chỉ parse / clang -fsyntax-only','Action Policy Engine chặn cấu hình sai mức áp/xung nhịp trước khi ghi BoardState'],
],widths=[3.20,4.55,4.25],size=15.5,row_h=.86,hdr_h=.48)
band(s,'ASR trên tập dev C7 đạt 0,0%; mã C sinh ra tuyệt đối không được thực thi trên máy chủ.',6.05,color='teal',size=19)
note(s,'35 giây. Trình bày 4 bề mặt tấn công chính, các chốt chặn xác định ở mức mã nguồn hiện có và kế hoạch P6.')

# 11 — Kế hoạch tới cuối kỳ (35 s)
s=base('7. Kế hoạch tuần 45–49 & Phạm vi dự phòng','Chia đều khối lượng P6–P9 cho 3 thành viên • Tự đánh giá: đúng tiến độ',source='Mốc bảo vệ vòng 1: Tuần 49 (02/12/2026) • plan/HANDOFF_P6_P9.md')
table(s,.65,1.60,12.0,['Tuần','Mục tiêu kỹ thuật & Sản phẩm bàn giao','Phân công (chia đều 3 thành viên)'],[
['Tuần 45 (02–08/11)','P6: Action Policy Engine, Taint Tracking, giảm over-refusal, mở rộng C7 v2 & cờ A7','Huy, Nhâm, Bình (phối hợp đồng đều)'],
['Tuần 46 (09–15/11)','P7: Hoàn thiện Project Memory xuyên phiên, Lesson Memory (Reflexion) & cờ A5','Nhâm, Bình, Huy (phối hợp đồng đều)'],
['Tuần 47 (16–22/11)','P8: Giao diện Web UI v2 (Plan, Board, Timeline, Interrupt panels) & quay Demo','Huy, Nhâm, Bình (phối hợp đồng đều)'],
['Tuần 48 (23–29/11)','P9: Bộ task độ bền C9, Fault Injection, khoá grader, chạy tập test & ablation A1–A7','Bình, Nhâm, Huy (phối hợp đồng đều)'],
['Tuần 49 (30/11–02/12)','Phân tích kiểm định H1–H9, hoàn thiện báo cáo D3/D4, slide và tổng duyệt bảo vệ','Cả 3 thành viên (33,33% mỗi người)'],
],widths=[2.65,6.15,3.20],size=16.5,row_h=.64)
txt(s,.80,5.68,11.7,.92,'• Giả định rủi ro nhất: lợi thế trên tập dev giữ vững có ý nghĩa thống kê trên tập test niêm phong.\n• Phạm vi thu hẹp dự phòng: giữ trọn bộ lõi v2 và tập trung kiểm chứng mô phỏng sâu cho STM32G474.',18,True,'amber')
note(s,'35 giây. Lộ trình tuần 45 đến tuần 49 được chia đều cho cả 3 thành viên, kèm phương án thu hẹp an toàn nếu cần.')

# 12 — Đóng góp & AI-use (25 s)
s=base('8. Phân công đóng góp đồng đều & AI-use Statement','Chia đều 33,33% cho mỗi thành viên ở cả D1, v2 (P0–P5) và kế hoạch P6–P9',source='Báo cáo tiến độ D2 — Mục 8')
table(s,.65,1.50,12.0,['Thành viên','Tỷ lệ','Vai trò trong D1, v2 (P0–P5) và Kế hoạch P6–P9'],[
['Huỳnh Đức Nhâm (2570276)','33,33%','• D1 & v2 (P0–P5): Kiến trúc, Agent Runtime v2 (state machine, planner, checkpoint), HITL Interrupt/Resume, Run API & ToolHub.\n• P6–P9: Project/Lesson Memory (P7), đồng thực hiện Policy Engine (P6), Web UI v2 (P8) và đánh giá cuối kỳ (P9).'],
['Hà Thanh Bình (2470732)','33,33%','• D1 & v2 (P0–P5): Pipeline dữ liệu P1 (SVD, AF, limits, hybrid index), text-manual/hw-knowledge, khung eval, 132 task & chạy S0–S3.\n• P6–P9: Task độ bền C9, Fault Injection, chạy tập test/ablation (P9) và đồng thực hiện P6–P8.'],
['La Quang Huy (2470736)','33,33%','• D1 & v2 (P0–P5): Threat Model, Guardrail an toàn & chống Injection, hwmodel (simulator, solver, DRC), board-state/hw-sandbox & Gate G1–G5.\n• P6–P9: Action Policy Engine & Taint Tracking (P6), Web UI v2 (P8) và đồng thực hiện P7, P9.'],
],widths=[2.65,1.10,8.25],size=14,row_h=.86,hdr_h=.42)
rect(s,.65,5.15,12.0,1.58,'pale',line='blue',radius=True)
txt(s,.85,5.23,11.6,1.42,'AI-use Statement:\n1. Thảo luận, phản biện và chốt phương án thiết kế kiến trúc, ý tưởng giải thuật cùng các thành viên nhóm.\n2. Triển khai bản kế hoạch kỹ thuật của các thành viên thành mã nguồn hiện thực và kịch bản kiểm thử.\n3. Xây dựng đề cương (outline) và hỗ trợ trình bày báo cáo tiến độ cùng bộ slide thuyết trình.',15.5,False,'ink')
note(s,'25 giây. Trình bày bảng đóng góp đồng đều 33,33% của 3 thành viên từ D1, v2 (P0–P5) đến P6–P9 và 3 hạng mục sử dụng AI.')

# 13 — Kết luận & Hỏi đáp (10 s)
s=base('Tổng kết tiến độ giữa kỳ D2 & Thảo luận','10 phút trình bày + Live Demo • 5 phút Hỏi & Đáp',source='Repository: github.com/htbitb/Agentic-RAG-for-Hardware-Specs-Smart-Retrieval-for-Embedded-IoT-System')
txt(s,.88,1.72,11.55,1.0,'Đã chuyển đổi thành công từ RAG tra cứu thụ động\nsang Hardware Co-Design Agent v2 có vòng phản hồi kín.',29,True,'blue')
txt(s,.90,3.15,11.45,2.3,'• Đã hoàn thành end-to-end P0–P5: Corpus thật (6 PDF, 3 SVD), 4 MCP Servers (32 tools),\n  BoardState + DRC, Simulator thanh ghi, Gate G1–G5 và Interrupt/Resume.\n• Đã đo trên 42 task dev × 3 seed: S3@P5 đạt TSR 87,3% (+31,0 điểm phần trăm so với S2D, p = 0,001).\n• Bước tiếp theo (P6–P9): Action Policy Engine, Lesson Memory, Web UI v2 và đánh giá trên tập test.',19.5)
band(s,'Xin cảm ơn Thầy và các bạn! Nhóm sẵn sàng nhận câu hỏi.',6.02,color='teal',size=20)
note(s,'10 giây. Kết thúc phần trình bày chính và chuyển sang 5 phút hỏi đáp. Các slide phụ lục phía sau hỗ trợ trả lời chi tiết.')

# 14 — Phụ lục: Phân tích lỗi còn lại
s=base('Phụ lục 1: Phân tích các trường hợp lỗi còn lại trên tập dev','16/126 run chính của S3@P5 chưa đạt • Ghi chú: Chưa test thật trên bo mạch điện',source='eval/results/dev_baseline_v2/report.md • scores.jsonl',backup=True)
table(s,.65,1.56,12.0,['Nhóm / Task','Số run trượt','Nguyên nhân kỹ thuật và hướng xử lý ở P6–P7'],[
['C4 (Đa lượt): C4-001, C4-011, C4-014','6 / 15 run','Agent giải được code cuối (final_state đạt 80%) nhưng đôi khi nêu cảnh báo xung đột chân PA8 / TIM3 / chân không FT muộn hơn lượt yêu cầu -> Cải thiện ở P6/P7.'],
['C6 (An toàn): C6-015','3 / 15 run','Guardrail v1 lọc từ khoá "5 V" từ chối quá mức (over-refusal) khi người dùng cấp 5 V vào chân VDDH hợp lệ của nRF52840 -> Xử lý bằng Policy Engine ở P6.'],
['C8 (Ngoài phạm vi): C8-008','2 / 9 run','Mô hình đưa ra ước lượng khoảng giá thương mại của linh kiện thay vì từ chối trả lời -> Bổ sung quy tắc kiểm soát câu hỏi ngoài tài liệu ở P6.'],
['C5 (Hệ thống): C5-002, C5-005, C5-011','5 / 12 run','Thiếu phép tính chi tiết ngưỡng VIH = 0,7 × VDD hoặc hết ngân sách bước khi chuỗi kiểm tra kéo dài -> Bổ sung Lesson Memory ở P7.'],
],widths=[3.15,1.45,7.40],row_h=.86,hdr_h=.48,size=15.5)
band(s,'Chưa test thật trên bo mạch điện; toàn bộ đánh giá hiện tại thực hiện qua simulator thanh ghi và kiểm tra cú pháp.',6.02,color='amber',size=18)
note(s,'Slide phụ lục giải thích chi tiết 16/126 run chưa đạt của S3@P5 và hướng khắc phục ở P6–P7.')

# 15 — Phụ lục: Trace tự sửa lỗi vật lý & nguồn trích dẫn
s=base('Phụ lục 2: Trace thực tế vòng tự sửa lỗi Clock & Provenance','S3@P5 • Task C3-020 • Seed 1 • Cấu hình PLL 168 MHz từ HSE 8 MHz trên STM32G474',source='eval/results/dev_baseline_v2/runs.jsonl:628 & scores.jsonl:628',backup=True)
table(s,.65,1.65,12.0,['Seq / Step','Sự kiện trong Trace (Run bd1fac3fa350)'],[
['54–55 / 24','simulate_register_writes lần 1 FAIL: SIM-CLK-LIMIT (tần số đầu vào PLL 2 MHz < mức tối thiểu 2,66 MHz)'],
['57–58 / 25','Agent tự tính lại hệ số chia M=1, N=42, R=2; mô phỏng lần 2 PASS nhưng còn cảnh báo trình tự chuyển clock'],
['62–63 / 27','Agent bổ sung bước chia AHB trung gian; mô phỏng lần 3 PASS hoàn toàn, SYSCLK = 168 MHz, violations = []'],
['82–83 / 31','Gọi finish lần 1 bị Gate từ chối: G1=fail (thiếu trích dẫn trang tài liệu cho dải PLL), G2–G5=pass'],
['95–97 / 32','Agent bổ sung trích dẫn [STM32G474_Datasheet p.123]; G1–G5 đều pass -> finish ACCEPTED'],
],widths=[2.35,9.65],size=17.5,row_h=.66)
band(s,'Kết quả grader độc lập: functional_sim PASS (168 MHz) và compile_check OK.',6.02,color='teal',size=19)
note(s,'Slide phụ lục minh hoạ vòng phản hồi kín: agent tự sửa lỗi vật lý SIM-CLK-LIMIT từ simulator và sửa lỗi trích dẫn G1 từ Verification Gate.')

# 16 — Phụ lục: Tái lập kết quả
s=base('Phụ lục 3: Hướng dẫn tái lập kết quả không tốn phí API','Repository công khai cùng bộ cache và manifest ghim mã băm SHA-256',source='eval/README.md • docs/reports/interim-report/interim-report.pdf',backup=True)
txt(s,.85,1.65,11.5,.42,'Lệnh phát lại (replay) từ bộ nhớ đệm, không gọi mạng:',21,True,'blue')
rect(s,.78,2.22,11.85,1.25,'pale',line='blue',radius=True)
txt(s,1.0,2.42,11.4,.85,'eval/.venv/bin/python -m eval.run --systems S3@P5\n  --suite smoke --seeds 1 --cache replay --out /tmp/interim-d2-replay',18,False,'ink',font='Courier New')
txt(s,.85,3.78,11.6,1.95,'• Chấm điểm lại toàn bộ 774 run (không tốn LLM):\n  eval/.venv/bin/python -m eval.score  --run eval/results/dev_baseline_v2\n  eval/.venv/bin/python -m eval.report --run eval/results/dev_baseline_v2\n• Kiểm chứng thực tế: 9/9 lượt chạy replay khớp từng byte với bản ghi gốc.',19.5)
band(s,'Repository: github.com/htbitb/Agentic-RAG-for-Hardware-Specs-Smart-Retrieval-for-Embedded-IoT-System',6.02,color='teal',size=17.5)
note(s,'Slide phụ lục cung cấp lệnh tái lập kết quả đánh giá và kiểm chứng từ cache.')

prs.core_properties.title='CO5151 — Báo cáo tiến độ D2 — Hardware Co-Design Agent v2'
prs.core_properties.subject='Báo cáo tiếng Việt: triển khai Agent v2 (S3@P5) và đánh giá tập dev'
prs.core_properties.author='Hà Thanh Bình; Huỳnh Đức Nhâm; La Quang Huy'
prs.core_properties.keywords='CO5151, D2, Agentic RAG, Hardware Co-Design Agent v2, tiếng Việt'
prs.core_properties.comments='13 slide chính (10 phút gồm live demo 2 phút) + 3 slide phụ lục.'
prs.save(OUT)
print(f'Wrote {OUT} ({len(prs.slides)} slides; 13 main + 3 backup)')
